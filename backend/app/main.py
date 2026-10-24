from __future__ import annotations
import os, re, io, json, sqlite3, time, uuid, sys
from pathlib import Path
from typing import Optional
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import httpx
from dotenv import load_dotenv
from .skills.registry import SKILLS, get_skill, public_skills

load_dotenv()
OLLAMA_BASE = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:7b")
DB_PATH = os.getenv("DATABASE_PATH", "./draftmate.db")
MAX_UPLOAD = int(os.getenv("MAX_UPLOAD_MB", "15")) * 1024 * 1024

app = FastAPI(title="DraftMate AI API", version="1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173","http://127.0.0.1:5173"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with db() as c:
        c.execute("""CREATE TABLE IF NOT EXISTS drafts (
          id TEXT PRIMARY KEY, title TEXT NOT NULL, skill TEXT NOT NULL, content TEXT NOT NULL,
          instruction TEXT NOT NULL, created_at INTEGER NOT NULL, updated_at INTEGER NOT NULL)""")
init_db()

class DraftRequest(BaseModel):
    title: str = Field(default="Untitled legal document", max_length=200)
    instruction: str = Field(min_length=2, max_length=12000)
    skill: str = "legal_notice"
    document: str = Field(default="", max_length=100000)
    references: str = Field(default="", max_length=60000)
    style_profile: str = Field(default="", max_length=8000)
    jurisdiction: str = "India"
    tone: str = "Formal and precise"

class ReviewRequest(BaseModel):
    document: str = Field(min_length=1, max_length=100000)

class SaveDraftRequest(BaseModel):
    id: Optional[str] = None
    title: str = Field(default="Untitled legal document", max_length=200)
    skill: str = "legal_notice"
    content: str = Field(default="", max_length=100000)
    instruction: str = Field(default="", max_length=12000)

def fallback(req: DraftRequest) -> str:
    label = get_skill(req.skill).name.upper()
    if req.skill == "summarize" and req.document.strip():
        return f"{label}\n\nSource text supplied:\n{req.document[:5000]}\n\nHUMAN REVIEW\nConfirm that this summary preserves all material obligations, dates, exceptions and defined terms."
    if req.skill in ("clause_review", "redline_review"):
        return f"""{label}

REQUEST
{req.instruction}

TEXT TO REVIEW
{req.document or "[Paste the relevant clause or document.]"} 

REVIEW CHECKLIST
1. Ambiguity and undefined terms
2. Conflicting obligations, dates, and conditions
3. Allocation of liability, remedies, and termination rights
4. Governing law, jurisdiction, notices, and dispute process
5. Missing facts and evidence

PROPOSED NEXT STEP
Confirm the commercial or legal objective and have the proposed wording reviewed by qualified counsel.

No legal authority has been independently verified."""
    return f"""{label}

Jurisdiction: {req.jurisdiction}
Instruction: {req.instruction}

1. PARTIES
[Insert full legal names, legal status, addresses and authority to sign.]

2. BACKGROUND AND PURPOSE
[Insert verified material facts and purpose. Do not state assumptions as facts.]

3. TERMS / MATERIAL FACTS
[Set out agreed facts, scope, obligations, dates, amounts and conditions.]

4. OBLIGATIONS AND TIMELINES
[Specify each responsible party, required action, deadline and evidence of completion.]

5. DISPUTES, GOVERNING LAW AND REMEDIES
[Insert provisions suitable for the transaction and jurisdiction after legal review.]

6. NOTICES
[Insert notice addresses and valid methods of service where appropriate.]

7. EXECUTION
[Insert signature blocks, date, place, witnesses and required formalities.]

OPEN ITEMS FOR HUMAN REVIEW
- Confirm all parties, facts, dates, amounts and supporting documents.
- Confirm applicable law, forum, limitation periods and procedural requirements.
- Verify every statute, case citation and quotation against authoritative sources.
- Replace all placeholders before use.

This is a starter template, not legal advice. No legal authorities have been independently verified."""

@app.get("/api/health")
async def health():
    connected = False
    models = []
    try:
        async with httpx.AsyncClient(timeout=2) as client:
            r = await client.get(f"{OLLAMA_BASE}/api/tags")
            if r.is_success:
                connected = True
                models = [x.get("name","") for x in r.json().get("models", [])]
    except Exception:
        pass
    return {"ok": True, "ollama_connected": connected, "model": OLLAMA_MODEL, "model_available": OLLAMA_MODEL in models, "available_models": models}

@app.get("/api/skills")
async def skills():
    return public_skills()

@app.get("/api/models")
async def models():
    """List models actually installed in the local Ollama instance."""
    try:
        async with httpx.AsyncClient(timeout=4) as client:
            r = await client.get(f"{OLLAMA_BASE}/api/tags")
            r.raise_for_status()
            names = [m.get("name", "") for m in r.json().get("models", []) if m.get("name")]
            return {"models": names, "default": OLLAMA_MODEL, "source": "local Ollama"}
    except Exception:
        return {"models": [], "default": OLLAMA_MODEL, "source": "Ollama unavailable", "warning": "Start Ollama to list installed models."}

@app.post("/api/prompt/improve")
async def improve_prompt(payload: dict):
    original = str(payload.get("instruction", "")).strip()[:12000]
    skill_id = str(payload.get("skill", "legal_notice"))
    jurisdiction = str(payload.get("jurisdiction", "India"))[:120]
    tone = str(payload.get("tone", "Formal and precise"))[:120]
    if len(original) < 2:
        raise HTTPException(400, "Enter a short instruction before improving it.")
    system = "Rewrite the user's legal drafting instruction into a clear task brief. Preserve the user's intent. Do not invent facts, dates, parties, legal authorities, amounts or deadlines. Use explicit placeholders for missing details. Return only the improved instruction, with a short list of needed details if important."
    prompt = f"Skill: {get_skill(skill_id).name}\nJurisdiction: {jurisdiction}\nTone: {tone}\nOriginal instruction: {original}\nMake it specific, structured and safe. Do not fill missing facts."
    try:
        async with httpx.AsyncClient(timeout=45) as client:
            r = await client.post(f"{OLLAMA_BASE}/api/chat", json={"model": OLLAMA_MODEL, "stream": False, "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}], "options": {"temperature": 0.15, "num_ctx": 4096}})
            r.raise_for_status()
            improved = r.json().get("message", {}).get("content", "").strip()
            if improved:
                return {"original": original, "improved": improved, "mode": "ollama", "model": OLLAMA_MODEL}
    except Exception:
        pass
    improved = (f"Prepare a {get_skill(skill_id).name.lower()} for the jurisdiction of {jurisdiction}.\n\n"
                f"OBJECTIVE AND FACTS\n{original}\n\n"
                f"DRAFTING REQUIREMENTS\n- Tone: {tone}.\n- Use only facts supplied in the instruction and reference materials.\n"
                "- Use clearly marked [PLACEHOLDER: ...] fields for missing parties, dates, amounts, deadlines or other material facts.\n"
                "- Do not invent case citations, statutory text, quotations or procedural requirements.\n"
                "- Structure the output with appropriate headings, numbered clauses and signature/execution blocks where relevant.\n"
                "- Finish with OPEN ITEMS FOR HUMAN REVIEW and identify authorities that require independent verification.")
    return {"original": original, "improved": improved, "mode": "local template", "warning": "Ollama was unavailable; this is a structured prompt template, not AI enhancement."}

@app.post("/api/draft")
async def generate_draft(req: DraftRequest):
    skill = get_skill(req.skill)
    system = f"""You are DraftMate AI, a cautious legal drafting assistant. Your current skill is {skill.name}: {skill.instruction}
Jurisdiction: {req.jurisdiction}. Tone: {req.tone}.
Never invent cases, statutes, citations, quotations, facts or procedural rules. Use [PLACEHOLDER] for missing facts. If no authoritative source is provided, state that legal authorities require verification. Preserve substantive meaning unless asked to change it. Distinguish sourced facts from assumptions. Include a concise OPEN ITEMS FOR HUMAN REVIEW section. You are not a lawyer and must not present output as definitive legal advice."""
    user = f"""TITLE: {req.title}
USER INSTRUCTION:
{req.instruction}

EXISTING DOCUMENT:
{req.document or "[No existing document]"}

REFERENCE MATERIAL:
{req.references or "[No references supplied]"}

STYLE PROFILE (user-provided examples only):
{req.style_profile or "[No style profile supplied]"}"""
    try:
        async with httpx.AsyncClient(timeout=150) as client:
            r = await client.post(f"{OLLAMA_BASE}/api/chat", json={"model": OLLAMA_MODEL, "stream": False, "messages":[{"role":"system","content":system},{"role":"user","content":user}], "options":{"temperature":0.2,"num_ctx":8192}})
            r.raise_for_status()
            content = r.json().get("message",{}).get("content","").strip()
            if not content: raise RuntimeError("Ollama returned an empty response")
            return {"content":content,"mode":"ollama","model":OLLAMA_MODEL,"warning":"Review all facts and authorities. The model does not independently verify citations."}
    except Exception as e:
        return {"content":fallback(req),"mode":"fallback","model":None,"warning":f"Template fallback used; Ollama unavailable or failed ({type(e).__name__}). Install/start Ollama for model generation."}

@app.post("/api/documents")
async def upload_document(file: UploadFile = File(...)):
    filename = file.filename or "reference"
    ext = Path(filename).suffix.lower()
    raw = await file.read()
    if len(raw) > MAX_UPLOAD: raise HTTPException(413, f"File exceeds configured {MAX_UPLOAD // (1024*1024)} MB limit.")
    try:
        if ext in (".txt",".md",".markdown",".csv"):
            text = raw.decode("utf-8",errors="replace")
        elif ext == ".pdf":
            from pypdf import PdfReader
            text = "\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        else: raise HTTPException(415,"Supported formats: PDF, DOCX, TXT, Markdown, CSV.")
    except HTTPException: raise
    except Exception as e: raise HTTPException(400,f"Unable to extract document: {e}")
    return {"filename":filename,"text":text[:60000],"characters":len(text),"truncated":len(text)>60000}

@app.post("/api/style-profile")
async def style_profile(payload: dict):
    text = str(payload.get("text",""))[:60000]
    if not text.strip(): raise HTTPException(400,"Provide reference text to analyze.")
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    headings = [x for x in lines if len(x)<100 and (x.isupper() or re.match(r"^\d+[\.\)]\s",x))]
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+",text) if s.strip()]
    avg = round(sum(len(s.split()) for s in sentences)/max(1,len(sentences)),1)
    formal = sum(bool(re.search(r"\b(hereby|whereas|pursuant|thereof|notwithstanding|shall)\b",s,re.I)) for s in sentences)
    profile = {"source":"user-provided reference","approx_words":len(text.split()),"average_sentence_words":avg,
               "heading_examples":headings[:12],"formal_legal_terms_detected":formal,
               "style_guidance":f"Mirror the reference's level of formality and heading structure. Average sentence length is about {avg} words. Use formal legal phrasing only where the source supports it. Preserve meaning; do not copy names, facts or clauses unless explicitly requested."}
    return profile

@app.post("/api/citations/review")
async def citation_review(req: ReviewRequest):
    patterns = [
      ("SCC Online",r"\b\d{4}\s+SCC\s+OnLine\s+[A-Za-z]+\s+\d+\b"),
      ("SCC",r"\b\d{4}\s+\(\d+\)\s+SCC\s+\d+\b"),
      ("AIR",r"\bAIR\s+\d{4}\s+[A-Z]+\s+\d+\b"),
      ("Section reference",r"\b(?:Section|Sec\.?)\s+\d+[A-Za-z]?(?:\([^)]+\))?\b"),
      ("Article reference",r"\bArticle\s+\d+[A-Za-z]?\b"),
      ("Rule reference",r"\bRule\s+\d+[A-Za-z]?\b")
    ]
    citations=[]
    for kind, pattern in patterns:
        for m in re.finditer(pattern,req.document,re.I):
            cite=m.group(0)
            if not any(x["citation"]==cite for x in citations):
                citations.append({"id":str(uuid.uuid4()),"citation":cite,"kind":kind,"status":"UNVERIFIED",
                    "evidence":None,"reason":"Citation-like text detected, but no authoritative source match has been performed."})
    sources=[
      {"name":"India Code","url":"https://www.indiacode.nic.in/","type":"official legislation portal","role":"Research starting point; not queried automatically by this prototype."},
      {"name":"Supreme Court of India","url":"https://www.sci.gov.in/","type":"official court website","role":"Research starting point; not queried automatically by this prototype."},
      {"name":"eCourts Services","url":"https://ecourts.gov.in/","type":"official court services","role":"Research starting point; not queried automatically by this prototype."},
      {"name":"Indian Kanoon","url":"https://indiankanoon.org/","type":"secondary research index","role":"Useful for discovery; independently confirm against authoritative sources."}
    ]
    return {"items":citations,"count":len(citations),"summary":{"verified":0,"partial":0,"unverified":len(citations)},
            "sources":sources,"method":"Pattern detection only; no source-level verification is configured. Therefore no citation can be labelled verified by this endpoint."}

@app.post("/api/drafts")
async def save_draft(req: SaveDraftRequest):
    now=int(time.time()); did=req.id or str(uuid.uuid4())
    with db() as c:
        c.execute("""INSERT INTO drafts(id,title,skill,content,instruction,created_at,updated_at) VALUES(?,?,?,?,?,?,?)
        ON CONFLICT(id) DO UPDATE SET title=excluded.title,skill=excluded.skill,content=excluded.content,instruction=excluded.instruction,updated_at=excluded.updated_at""",
        (did,req.title,req.skill,req.content,req.instruction,now,now))
    return {"id":did,"title":req.title,"saved":True,"updated_at":now}

@app.get("/api/drafts")
async def list_drafts():
    with db() as c:
        rows=c.execute("SELECT id,title,skill,created_at,updated_at FROM drafts ORDER BY updated_at DESC LIMIT 50").fetchall()
    return [dict(r) for r in rows]

@app.get("/api/drafts/{draft_id}")
async def get_draft(draft_id: str):
    with db() as c: row=c.execute("SELECT * FROM drafts WHERE id=?",(draft_id,)).fetchone()
    if not row: raise HTTPException(404,"Draft not found")
    return dict(row)

@app.delete("/api/drafts/{draft_id}")
async def delete_draft(draft_id: str):
    with db() as c: cur=c.execute("DELETE FROM drafts WHERE id=?",(draft_id,))
    if cur.rowcount==0: raise HTTPException(404,"Draft not found")
    return {"deleted":True}

@app.post("/api/export/docx")
async def export_docx(payload: dict):
    from docx import Document
    title=str(payload.get("title","DraftMate document"))[:200]
    content=str(payload.get("content",""))[:100000]
    doc=Document(); doc.add_heading(title,0)
    for line in content.splitlines():
        if line.startswith("# "): doc.add_heading(line[2:],level=1)
        elif line.startswith("## "): doc.add_heading(line[3:],level=2)
        elif line.startswith("### "): doc.add_heading(line[4:],level=3)
        elif line.strip(): doc.add_paragraph(line)
        else: doc.add_paragraph("")
    buf=io.BytesIO(); doc.save(buf)
    safe=re.sub(r"[^A-Za-z0-9_-]+","_",title).strip("_") or "draft"
    return Response(buf.getvalue(),media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":f'attachment; filename="{safe[:80]}.docx"'})


# Serve the built React frontend when running as the desktop application.
# API routes above take precedence; this mount is deliberately added last.
_default_static = (Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2])) / "static")
_static_dir = Path(os.getenv("DRAFTMATE_STATIC_DIR", str(_default_static))).resolve()
if _static_dir.is_dir():
    app.mount("/", StaticFiles(directory=str(_static_dir), html=True), name="frontend")
